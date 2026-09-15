from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE_DIR = Path(r"C:\Users\celvn\Downloads")
OUTPUT_DIR = Path(__file__).resolve().parents[1] / "docs" / "revised_questionnaires"


INTRO_EN = (
    "English. This questionnaire evaluates only the GradTrack functions implemented and accessible to your assigned "
    "user role. Answer each statement based on your actual use during system evaluation. For a conditional function, "
    "answer only when the stated condition applies. Your responses will be used only for academic purposes and will "
    "be treated confidentially."
)
INTRO_TL = (
    "Tagalog. Sinusuri ng talatanungang ito ang mga GradTrack function na aktuwal na ipinatupad at maaaring ma-access "
    "ng inyong nakatalagang user role. Sagutin ang bawat pahayag batay sa inyong aktuwal na paggamit sa system evaluation. "
    "Para sa conditional na function, sumagot lamang kapag naaangkop ang nakasaad na kondisyon. Ang inyong mga sagot ay "
    "gagamitin lamang para sa layuning akademiko at pananatilihing kumpidensyal."
)


RESEARCH_GROUPS = [
    [
        ("I can access the GradTrack Dashboard after signing in with my Research Coordinator account.",
         "Maaari kong ma-access ang GradTrack Dashboard pagkatapos mag-sign in gamit ang aking Research Coordinator account."),
        ("I can view the overall Employment Rate and its employed-versus-valid-respondent counts.",
         "Maaari kong makita ang overall Employment Rate at ang bilang ng employed kumpara sa valid respondents."),
        ("I can view the overall Alignment Rate and its aligned-versus-valid-applicable-response counts.",
         "Maaari kong makita ang overall Alignment Rate at ang bilang ng aligned kumpara sa valid applicable responses."),
        ("I can view the number of active surveys and the title of the survey represented on the dashboard.",
         "Maaari kong makita ang bilang ng active surveys at ang pamagat ng survey na kinakatawan sa dashboard."),
    ],
    [
        ("I can view the survey coverage percentage for the survey represented on the dashboard.",
         "Maaari kong makita ang survey coverage percentage para sa survey na kinakatawan sa dashboard."),
        ("I can view the number of valid responses and eligible active graduates.",
         "Maaari kong makita ang bilang ng valid responses at eligible active graduates."),
        ("I can view the number of graduates with pending survey responses.",
         "Maaari kong makita ang bilang ng mga graduate na may pending survey response."),
        ("I can compare valid and pending response counts using the survey participation summary.",
         "Maaari kong maikumpara ang valid at pending response counts gamit ang survey participation summary."),
    ],
    [
        ("I can view the Employability Index by Program bar chart.",
         "Maaari kong makita ang Employability Index by Program bar chart."),
        ("I can identify each program represented by its program code in the employability chart.",
         "Maaari kong matukoy ang bawat programa sa employability chart gamit ang program code nito."),
        ("I can view the employability percentage displayed for each program.",
         "Maaari kong makita ang employability percentage na ipinapakita para sa bawat programa."),
        ("I can expand a program breakdown to view its employed and valid-response counts.",
         "Maaari kong i-expand ang program breakdown upang makita ang employed at valid-response counts nito."),
    ],
    [
        ("I can view the Employment Trends line chart by graduation year.",
         "Maaari kong makita ang Employment Trends line chart ayon sa graduation year."),
        ("I can view employment-rate values across the graduation years shown.",
         "Maaari kong makita ang employment-rate values sa mga graduation year na ipinapakita."),
        ("I can view alignment-rate values across the graduation years shown.",
         "Maaari kong makita ang alignment-rate values sa mga graduation year na ipinapakita."),
        ("I can compare the employment-rate and alignment-rate trend lines.",
         "Maaari kong maikumpara ang employment-rate at alignment-rate trend lines."),
    ],
    [
        ("I can view the overall Job Alignment Distribution.",
         "Maaari kong makita ang overall Job Alignment Distribution."),
        ("I can filter the Job Alignment Distribution by program.",
         "Maaari kong i-filter ang Job Alignment Distribution ayon sa programa."),
        ("I can view the counts and percentages for the displayed alignment categories.",
         "Maaari kong makita ang bilang at percentage para sa mga ipinapakitang alignment category."),
        ("I can compare the overall alignment distribution with a selected program's distribution.",
         "Maaari kong maikumpara ang overall alignment distribution sa distribution ng napiling programa."),
    ],
    [
        ("I can view the title of the survey represented in the Selected Survey Snapshot.",
         "Maaari kong makita ang pamagat ng survey na kinakatawan sa Selected Survey Snapshot."),
        ("I can view the survey coverage bar and percentage in the snapshot.",
         "Maaari kong makita ang survey coverage bar at percentage sa snapshot."),
        ("I can view valid-response and pending-response counts in the snapshot.",
         "Maaari kong makita ang valid-response at pending-response counts sa snapshot."),
        ("I can view employed and unemployed respondent counts in the snapshot.",
         "Maaari kong makita ang bilang ng employed at unemployed respondents sa snapshot."),
    ],
    [
        ("I can view aligned-job and not-aligned-job counts in the survey snapshot.",
         "Maaari kong makita ang bilang ng aligned-job at not-aligned-job responses sa survey snapshot."),
        ("I can distinguish an unavailable rate from a zero-percent rate when no valid responses exist.",
         "Maaari kong matukoy ang unavailable rate kumpara sa zero-percent rate kapag walang valid responses."),
        ("I can identify the valid respondent base used for the displayed Employment Rate.",
         "Maaari kong matukoy ang valid respondent base na ginamit para sa ipinapakitang Employment Rate."),
        ("I can identify the valid applicable-response base used for the displayed Alignment Rate.",
         "Maaari kong matukoy ang valid applicable-response base na ginamit para sa ipinapakitang Alignment Rate."),
    ],
    [
        ("I can use chart tooltips to read values in the employability chart.",
         "Maaari kong gamitin ang chart tooltips upang basahin ang values sa employability chart."),
        ("I can use chart tooltips to read values in the employment-trends chart.",
         "Maaari kong gamitin ang chart tooltips upang basahin ang values sa employment-trends chart."),
        ("I can use the pie-chart legend to identify each job-alignment category.",
         "Maaari kong gamitin ang pie-chart legend upang matukoy ang bawat job-alignment category."),
        ("I can recognize when the dashboard reports that no program, trend, or alignment data are available.",
         "Maaari kong matukoy kapag ipinapakita ng dashboard na walang available na program, trend, o alignment data."),
    ],
    [
        ("I can distinguish programs through the color-coded bars in the employability chart.",
         "Maaari kong matukoy ang iba't ibang programa gamit ang color-coded bars sa employability chart."),
        ("I can read the graduation-year and percentage axes used in the employment-trends chart.",
         "Maaari kong basahin ang graduation-year at percentage axes na ginagamit sa employment-trends chart."),
        ("I can use the chart legend to distinguish Employment Rate from Alignment Rate.",
         "Maaari kong gamitin ang chart legend upang maihiwalay ang Employment Rate sa Alignment Rate."),
        ("I can see the alignment pie chart and legend update after choosing a program filter.",
         "Maaari kong makita ang pag-update ng alignment pie chart at legend pagkatapos pumili ng program filter."),
    ],
    [
        ("I can review graduate tracer analytics without controls for changing survey or graduate records.",
         "Maaari kong suriin ang graduate tracer analytics nang walang controls para baguhin ang survey o graduate records."),
        ("I can open My Profile from the account menu.",
         "Maaari kong buksan ang My Profile mula sa account menu."),
        ("I can update my account name, profile image, and password from My Profile.",
         "Maaari kong i-update ang aking account name, profile image, at password mula sa My Profile."),
        ("I can use the theme control and, when enabled, the notification control in the dashboard header.",
         "Maaari kong gamitin ang theme control at, kapag enabled, ang notification control sa dashboard header."),
    ],
]


DEAN_GROUPS = [
    [
        ("I can access the Graduate Survey Participation page for my assigned program scope.",
         "Maaari kong ma-access ang Graduate Survey Participation page para sa aking nakatalagang program scope."),
        ("I can see the program code or codes included in my assigned scope.",
         "Maaari kong makita ang program code o mga code na kasama sa aking nakatalagang scope."),
        ("I can view the total number of graduates in my assigned scope.",
         "Maaari kong makita ang kabuuang bilang ng mga graduate sa aking nakatalagang scope."),
        ("I can view the number of graduates who answered the active survey.",
         "Maaari kong makita ang bilang ng mga graduate na sumagot sa active survey."),
        ("I can view the number of graduates with no response to the active survey.",
         "Maaari kong makita ang bilang ng mga graduate na walang response sa active survey."),
        ("I can compare the Total, Answered, and No Survey Response summary cards.",
         "Maaari kong maikumpara ang Total, Answered, at No Survey Response summary cards."),
        ("I can view participation records only for programs assigned to my Dean account.",
         "Maaari kong makita ang participation records para lamang sa mga programang nakatalaga sa aking Dean account."),
        ("I can identify the program and graduation year of each graduate in the list.",
         "Maaari kong matukoy ang programa at graduation year ng bawat graduate sa listahan."),
    ],
    [
        ("I can search participation records by graduate name.",
         "Maaari akong maghanap sa participation records gamit ang pangalan ng graduate."),
        ("I can search participation records by email address.",
         "Maaari akong maghanap sa participation records gamit ang email address."),
        ("I can search participation records by student ID.",
         "Maaari akong maghanap sa participation records gamit ang student ID."),
        ("I can filter records as All Status, Answered, or Not Answered.",
         "Maaari kong i-filter ang records bilang All Status, Answered, o Not Answered."),
        ("I can filter participation records by graduation year.",
         "Maaari kong i-filter ang participation records ayon sa graduation year."),
        ("I can combine the search, response-status, and graduation-year filters.",
         "Maaari kong pagsabayin ang search, response-status, at graduation-year filters."),
        ("I can move between pages when the matching participation records exceed one page.",
         "Maaari akong lumipat sa pagitan ng mga page kapag lumampas sa isang page ang matching participation records."),
        ("I can see which active survey is being used for participation monitoring.",
         "Maaari kong makita kung aling active survey ang ginagamit para sa participation monitoring."),
    ],
    [
        ("I can view each graduate's student ID in the participation list.",
         "Maaari kong makita ang student ID ng bawat graduate sa participation list."),
        ("I can view each graduate's name and email address in the participation list.",
         "Maaari kong makita ang pangalan at email address ng bawat graduate sa participation list."),
        ("I can identify whether an individual graduate is Answered or Not Answered.",
         "Maaari kong matukoy kung ang isang graduate ay Answered o Not Answered."),
        ("I can identify graduates whose records do not contain an email address.",
         "Maaari kong matukoy ang mga graduate na walang email address sa kanilang record."),
        ("I can select an unanswered graduate whose record contains an email address.",
         "Maaari kong piliin ang isang unanswered graduate na may email address sa record."),
        ("I can select all eligible unanswered graduates shown on the current page.",
         "Maaari kong piliin ang lahat ng eligible unanswered graduates na nasa kasalukuyang page."),
        ("I can clear selected graduates by using the page-selection checkbox again.",
         "Maaari kong alisin ang pagkakapili sa mga graduate sa pamamagitan ng muling paggamit ng page-selection checkbox."),
        ("I can select all matching graduates with no response across the current search and year filters.",
         "Maaari kong piliin ang lahat ng matching graduates na walang response ayon sa kasalukuyang search at year filters."),
    ],
    [
        ("I can open View Answers for a graduate who submitted the active survey.",
         "Maaari kong buksan ang View Answers para sa graduate na nagsumite ng active survey."),
        ("I cannot open survey answers for a graduate who has not submitted a response.",
         "Hindi ko mabubuksan ang survey answers para sa graduate na hindi pa nagsumite ng response."),
        ("I can identify the graduate and survey represented in the answer viewer.",
         "Maaari kong matukoy ang graduate at survey na kinakatawan sa answer viewer."),
        ("I can view each submitted question and its recorded answer.",
         "Maaari kong makita ang bawat submitted question at ang recorded answer nito."),
        ("I can view submitted answers grouped under their survey sections.",
         "Maaari kong makita ang submitted answers na nakapangkat ayon sa survey sections."),
        ("I can view structured answers when a response contains multiple values.",
         "Maaari kong makita ang structured answers kapag maraming value ang isang response."),
        ("I can close the answer viewer and return to the participation list.",
         "Maaari kong isara ang answer viewer at bumalik sa participation list."),
        ("I can review submitted answers without changing the graduate's response.",
         "Maaari kong suriin ang submitted answers nang hindi binabago ang response ng graduate."),
    ],
    [
        ("I can select graduates with no survey response before sending reminders.",
         "Maaari kong piliin ang mga graduate na walang survey response bago magpadala ng reminders."),
        ("I can edit the email message used for the reminder.",
         "Maaari kong i-edit ang email message na gagamitin sa reminder."),
        ("I can see the number of graduates currently selected for notification.",
         "Maaari kong makita ang bilang ng mga graduate na kasalukuyang napili para sa notification."),
        ("I can send reminders only while an active survey is available.",
         "Maaari akong magpadala ng reminders kapag may available na active survey."),
        ("I receive a confirmation prompt before reminder emails are sent.",
         "Nakakatanggap ako ng confirmation prompt bago ipadala ang reminder emails."),
        ("The reminder function sends only to unanswered graduates and skips missing or invalid email addresses.",
         "Nagpapadala lamang ang reminder function sa unanswered graduates at nilalampasan ang missing o invalid email addresses."),
        ("I can view the eligible, sent, skipped, and failed counts after a reminder run.",
         "Maaari kong makita ang eligible, sent, skipped, at failed counts pagkatapos ng reminder run."),
        ("I can view refreshed participation information after reminders are processed.",
         "Maaari kong makita ang refreshed participation information pagkatapos maproseso ang reminders."),
    ],
]


REGISTRAR_GROUPS = [
    [
        ("I can add an individual graduate from the Manage Graduates page.",
         "Maaari akong magdagdag ng indibidwal na graduate mula sa Manage Graduates page."),
        ("I can enter the graduate's student number and complete name, including a name extension when applicable.",
         "Maaari kong ilagay ang student number at kumpletong pangalan ng graduate, kasama ang name extension kapag naaangkop."),
        ("I can enter the graduate's email address and contact number.",
         "Maaari kong ilagay ang email address at contact number ng graduate."),
        ("I can assign the graduate's program and year graduated.",
         "Maaari kong italaga ang programa at year graduated ng graduate."),
        ("The form checks the required student-number format before saving.",
         "Sinusuri ng form ang kinakailangang student-number format bago mag-save."),
        ("The form checks the required mobile-number format before saving.",
         "Sinusuri ng form ang kinakailangang mobile-number format bago mag-save."),
        ("The form can infer a graduation year from the student number and selected program when the available data allow it.",
         "Maaaring matukoy ng form ang graduation year mula sa student number at napiling programa kapag sapat ang available na data."),
        ("I receive a clear success or error message after attempting to add a graduate.",
         "Nakakatanggap ako ng malinaw na success o error message pagkatapos subukang magdagdag ng graduate."),
    ],
    [
        ("I can import graduate records from an XLSX or CSV file.",
         "Maaari akong mag-import ng graduate records mula sa XLSX o CSV file."),
        ("The import can read a labeled header row from the spreadsheet.",
         "Mababasa ng import ang labeled header row mula sa spreadsheet."),
        ("The import can map student, name, contact, program, and graduation-year data from recognized columns.",
         "Maaaring i-map ng import ang student, name, contact, program, at graduation-year data mula sa recognized columns."),
        ("The import can use the selected department when a row has no recognized program value.",
         "Maaaring gamitin ng import ang napiling department kapag walang recognized program value ang isang row."),
        ("The import can use the official-list heading or selected year when a row needs a graduation year.",
         "Maaaring gamitin ng import ang official-list heading o napiling year kapag kailangan ng graduation year ng isang row."),
        ("The import skips rows that do not contain the required name or graduation-year data.",
         "Nilalampasan ng import ang rows na walang kinakailangang name o graduation-year data."),
        ("I can view how many imported records were added and how many failed.",
         "Maaari kong makita kung ilang imported records ang naidagdag at kung ilan ang nag-fail."),
        ("I can view summarized reasons for rows skipped during import.",
         "Maaari kong makita ang summarized reasons para sa rows na nilaktawan sa import."),
    ],
    [
        ("I can search graduate records by name.",
         "Maaari akong maghanap ng graduate records gamit ang pangalan."),
        ("I can search graduate records by student ID.",
         "Maaari akong maghanap ng graduate records gamit ang student ID."),
        ("I can search graduate records by email address.",
         "Maaari akong maghanap ng graduate records gamit ang email address."),
        ("I can filter graduate records by department.",
         "Maaari kong i-filter ang graduate records ayon sa department."),
        ("I can filter graduate records by year graduated.",
         "Maaari kong i-filter ang graduate records ayon sa year graduated."),
        ("I can combine search, department, and graduation-year filters.",
         "Maaari kong pagsabayin ang search, department, at graduation-year filters."),
        ("I can view the number of active or archived records in the selected list.",
         "Maaari kong makita ang bilang ng active o archived records sa napiling listahan."),
        ("I can move between result pages when matching records exceed one page.",
         "Maaari akong lumipat sa pagitan ng result pages kapag lumampas sa isang page ang matching records."),
    ],
    [
        ("I can select an individual active graduate record for archiving.",
         "Maaari kong piliin ang isang active graduate record para i-archive."),
        ("I can select all graduate records shown on the current page.",
         "Maaari kong piliin ang lahat ng graduate records na nasa kasalukuyang page."),
        ("I can archive multiple selected graduate records in one action.",
         "Maaari kong i-archive ang maraming napiling graduate records sa isang action."),
        ("I can archive all records for a selected graduation year within the selected department.",
         "Maaari kong i-archive ang lahat ng records para sa napiling graduation year sa loob ng napiling department."),
        ("I receive a confirmation prompt before records are archived.",
         "Nakakatanggap ako ng confirmation prompt bago i-archive ang records."),
        ("The confirmation explains that related accounts, survey responses, profiles, forum history, messages, and job records are preserved.",
         "Ipinapaliwanag ng confirmation na mapapanatili ang related accounts, survey responses, profiles, forum history, messages, at job records."),
        ("Archived records leave the active Manage Graduates list.",
         "Nawawala sa active Manage Graduates list ang archived records."),
        ("I can open the Registrar Archive to manage archived graduate records.",
         "Maaari kong buksan ang Registrar Archive upang pamahalaan ang archived graduate records."),
    ],
    [
        ("I can search and filter records while viewing the Registrar Archive.",
         "Maaari akong mag-search at mag-filter ng records habang nasa Registrar Archive."),
        ("I can view when an archived record was archived.",
         "Maaari kong makita kung kailan na-archive ang isang record."),
        ("I can view which administrator archived a record.",
         "Maaari kong makita kung sinong administrator ang nag-archive ng isang record."),
        ("I can restore an individual archived graduate record.",
         "Maaari kong i-restore ang isang archived graduate record."),
        ("A restored record returns to active graduate management with its relationships intact.",
         "Bumabalik sa active graduate management ang restored record nang buo ang mga kaugnay nitong data."),
        ("I can permanently delete an individual record only from the Registrar Archive.",
         "Maaari kong permanenteng i-delete ang isang record mula lamang sa Registrar Archive."),
        ("I can permanently delete multiple selected archived records.",
         "Maaari kong permanenteng i-delete ang maraming napiling archived records."),
        ("I receive an irreversible-action warning before permanent deletion.",
         "Nakakatanggap ako ng warning na hindi na maibabalik ang action bago ang permanent deletion."),
    ],
]


SUPER_ADMIN_GROUPS = [
    [
        ("I can search staff accounts and filter them by role or active status.",
         "Maaari akong maghanap ng staff accounts at i-filter ang mga ito ayon sa role o active status."),
        ("I can create a staff account with a name, username, email, role, password, and status.",
         "Maaari akong gumawa ng staff account na may name, username, email, role, password, at status."),
        ("I can edit an existing staff account's identity and assigned role.",
         "Maaari kong i-edit ang identity at assigned role ng isang existing staff account."),
        ("I can set or change a staff account password subject to the password rules.",
         "Maaari akong magtakda o magpalit ng password ng staff account ayon sa password rules."),
        ("I can activate or deactivate another staff account, while my signed-in Super Admin account is protected from self-deactivation.",
         "Maaari kong i-activate o i-deactivate ang ibang staff account, habang protektado ang naka-sign-in kong Super Admin account laban sa self-deactivation."),
    ],
    [
        ("I can view the active survey, eligible-graduate count, reminder frequency, and total reminders sent.",
         "Maaari kong makita ang active survey, eligible-graduate count, reminder frequency, at total reminders sent."),
        ("I can view the graduates currently eligible to receive survey reminders.",
         "Maaari kong makita ang mga graduate na kasalukuyang eligible makatanggap ng survey reminders."),
        ("I can customize the subject and message before sending manual reminder emails.",
         "Maaari kong i-customize ang subject at message bago magpadala ng manual reminder emails."),
        ("I can configure a preset or custom reminder interval for the external reminder scheduler.",
         "Maaari akong mag-configure ng preset o custom reminder interval para sa external reminder scheduler."),
        ("I can review reminder history with sent, failed, and skipped results.",
         "Maaari kong suriin ang reminder history na may sent, failed, at skipped results."),
    ],
    [
        ("I can search and filter audit logs by role, department, action, module, and date range.",
         "Maaari akong mag-search at mag-filter ng audit logs ayon sa role, department, action, module, at date range."),
        ("I can open an audit entry to view its description, previous values, new values, and additional metadata.",
         "Maaari kong buksan ang audit entry upang makita ang description, previous values, new values, at additional metadata."),
        ("I can export the currently filtered audit trail as a CSV file.",
         "Maaari kong i-export bilang CSV file ang kasalukuyang filtered audit trail."),
        ("I can view the database name, table count, record count, approximate size, and included tables before backup.",
         "Maaari kong makita ang database name, table count, record count, approximate size, at included tables bago mag-backup."),
        ("I can download a restorable SQL backup of the GradTrack database.",
         "Maaari akong mag-download ng restorable SQL backup ng GradTrack database."),
    ],
    [
        ("I can edit the system name, institution details, public contact information, and footer text.",
         "Maaari kong i-edit ang system name, institution details, public contact information, at footer text."),
        ("I can upload or change the system logo, login logo, favicon, and login background image.",
         "Maaari akong mag-upload o magpalit ng system logo, login logo, favicon, at login background image."),
        ("I can change the primary and secondary theme colors.",
         "Maaari kong baguhin ang primary at secondary theme colors."),
        ("I can edit the configurable title, welcome text, subtitle, and additional text on the login page.",
         "Maaari kong i-edit ang configurable title, welcome text, subtitle, at additional text sa login page."),
        ("I can preview the configured login-page appearance before saving it.",
         "Maaari kong i-preview ang configured login-page appearance bago ito i-save."),
    ],
    [
        ("I can enable or disable the graduate survey, job support, community forum, messaging, and notification modules.",
         "Maaari kong i-enable o i-disable ang graduate survey, job support, community forum, messaging, at notification modules."),
        ("I can edit the survey title and instructions shown before graduate verification.",
         "Maaari kong i-edit ang survey title at instructions na ipinapakita bago ang graduate verification."),
        ("I can configure survey availability and the messages shown when the survey is unavailable or completed.",
         "Maaari kong i-configure ang survey availability at ang messages na ipinapakita kapag unavailable o completed ang survey."),
        ("I can edit the community guidelines and default community announcement.",
         "Maaari kong i-edit ang community guidelines at default community announcement."),
        ("I can enable or disable media uploads for community forum posts.",
         "Maaari kong i-enable o i-disable ang media uploads para sa community forum posts."),
    ],
    [
        ("I can view whether Maintenance Mode is currently on or off.",
         "Maaari kong makita kung kasalukuyang on o off ang Maintenance Mode."),
        ("I can turn Maintenance Mode on after confirming the action.",
         "Maaari kong i-on ang Maintenance Mode pagkatapos kumpirmahin ang action."),
        ("I can turn Maintenance Mode off when normal access should resume.",
         "Maaari kong i-off ang Maintenance Mode kapag dapat nang ibalik ang normal access."),
        ("I can edit the maintenance-page title, main message, and expected-availability message.",
         "Maaari kong i-edit ang maintenance-page title, main message, at expected-availability message."),
        ("Maintenance Mode blocks regular users while preserving Super Admin access.",
         "Bina-block ng Maintenance Mode ang regular users habang pinananatili ang Super Admin access."),
    ],
    [
        ("I can edit the text and images used in the About page sections.",
         "Maaari kong i-edit ang text at images na ginagamit sa About page sections."),
        ("I can create, edit, reorder, show, hide, or delete FAQ categories and questions.",
         "Maaari akong gumawa, mag-edit, mag-reorder, magpakita, magtago, o mag-delete ng FAQ categories at questions."),
        ("I can edit the Privacy Policy introduction, dates, section headings, and formatted content.",
         "Maaari kong i-edit ang Privacy Policy introduction, dates, section headings, at formatted content."),
        ("I can add, reorder, show, hide, or delete Privacy Policy sections.",
         "Maaari akong magdagdag, mag-reorder, magpakita, magtago, o mag-delete ng Privacy Policy sections."),
        ("I can preview unsaved About, FAQ, and Privacy Policy changes before publishing them.",
         "Maaari kong i-preview ang unsaved About, FAQ, at Privacy Policy changes bago i-publish ang mga ito."),
    ],
    [
        ("I can save changes made in each System Settings tab.",
         "Maaari kong i-save ang mga pagbabagong ginawa sa bawat System Settings tab."),
        ("I can reset an individual setting to its default value before saving.",
         "Maaari kong i-reset ang isang setting sa default value nito bago mag-save."),
        ("I receive validation feedback when a setting value, uploaded image, email address, or color contrast is unacceptable.",
         "Nakakatanggap ako ng validation feedback kapag hindi katanggap-tanggap ang setting value, uploaded image, email address, o color contrast."),
        ("I can update my own Super Admin profile name, profile image, and password.",
         "Maaari kong i-update ang name, profile image, at password ng sarili kong Super Admin profile."),
        ("I can navigate separately to User Management, Auto Email Reminders, Audit Trail, Backup Database, and System Settings.",
         "Maaari akong mag-navigate nang hiwalay sa User Management, Auto Email Reminders, Audit Trail, Backup Database, at System Settings."),
    ],
]


ALUMNI_GRADUATE_GROUPS = [
    [
        ("I can verify my identity for an active survey using either my student number or email address together with my last name and program.",
         "Maaari kong i-verify ang aking identity para sa active survey gamit ang student number o email address kasama ang aking last name at programa."),
        ("The system checks whether my graduate record and graduation year are eligible for the active survey.",
         "Sinusuri ng system kung eligible ang aking graduate record at graduation year para sa active survey."),
        ("I can review and accept the Data Privacy Notice and Informed Consent before answering.",
         "Maaari kong basahin at tanggapin ang Data Privacy Notice and Informed Consent bago sumagot."),
        ("I can review and edit survey fields prefilled from my verified registrar record.",
         "Maaari kong suriin at i-edit ang survey fields na na-prefill mula sa aking verified registrar record."),
        ("I can answer the question types presented in the active survey, including applicable text, choice, checkbox, and date fields.",
         "Maaari kong sagutan ang mga question type sa active survey, kabilang ang naaangkop na text, choice, checkbox, at date fields."),
        ("I can move between survey sections, while required-field validation identifies incomplete answers.",
         "Maaari akong lumipat sa pagitan ng survey sections habang tinutukoy ng required-field validation ang mga kulang na sagot."),
        ("My unfinished survey answers are saved as a draft that I can resume in the same browser.",
         "Nase-save bilang draft ang aking unfinished survey answers upang maipagpatuloy ko ang mga ito sa parehong browser."),
        ("After submitting the survey, I can request a Graduate Portal account using the prefilled details and a password, subject to Alumni Admin approval.",
         "Pagkatapos magsumite ng survey, maaari akong humiling ng Graduate Portal account gamit ang prefilled details at password, na sasailalim sa Alumni Admin approval."),
    ],
    [
        ("I can read announcements and open their complete details in the Graduate Portal.",
         "Maaari kong basahin ang announcements at buksan ang kumpletong details nito sa Graduate Portal."),
        ("I can view dashboard counts for published forum posts, my posts, conversations, and approved jobs.",
         "Maaari kong makita sa dashboard ang counts para sa published forum posts, aking posts, conversations, at approved jobs."),
        ("I can view my alumni score, employment status, course-alignment status, and earned badges.",
         "Maaari kong makita ang aking alumni score, employment status, course-alignment status, at earned badges."),
        ("I can search approved jobs by title, company, skills, location, or program fit.",
         "Maaari akong maghanap ng approved jobs ayon sa title, company, skills, location, o program fit."),
        ("I can open an approved job to view its details and available application instructions.",
         "Maaari kong buksan ang approved job upang makita ang details at available application instructions nito."),
        ("I can access job-post creation only when my employment status is marked employed.",
         "Maaari kong ma-access ang job-post creation kapag marked employed ang aking employment status."),
        ("When eligible, I can create, edit, and delete my own job posts.",
         "Kapag eligible, maaari akong gumawa, mag-edit, at mag-delete ng sarili kong job posts."),
        ("I can view the approval status and review notes for my submitted job posts.",
         "Maaari kong makita ang approval status at review notes para sa aking submitted job posts."),
    ],
    [
        ("I can search community posts and filter them by category, program, and graduation year.",
         "Maaari akong maghanap ng community posts at i-filter ang mga ito ayon sa category, program, at graduation year."),
        ("I can create a community post and, when media uploads are enabled, attach supported images or videos.",
         "Maaari akong gumawa ng community post at, kapag enabled ang media uploads, mag-attach ng supported images o videos."),
        ("I can edit or delete my own community post.",
         "Maaari kong i-edit o i-delete ang sarili kong community post."),
        ("I can like or unlike a community post.",
         "Maaari kong i-like o i-unlike ang isang community post."),
        ("I can open a post discussion to read and add comments.",
         "Maaari kong buksan ang post discussion upang magbasa at magdagdag ng comments."),
        ("I can report another user's post or comment and provide a reason.",
         "Maaari kong i-report ang post o comment ng ibang user at magbigay ng reason."),
        ("I can open the community profile of another graduate from forum content.",
         "Maaari kong buksan ang community profile ng ibang graduate mula sa forum content."),
        ("I can view the published or hidden status of my own forum posts.",
         "Maaari kong makita ang published o hidden status ng sarili kong forum posts."),
    ],
    [
        ("I can search my direct and group conversations.",
         "Maaari akong maghanap sa aking direct at group conversations."),
        ("I can start a direct conversation with another graduate.",
         "Maaari akong magsimula ng direct conversation sa ibang graduate."),
        ("I can create a group chat by selecting at least two other graduates.",
         "Maaari akong gumawa ng group chat sa pamamagitan ng pagpili ng hindi bababa sa dalawang ibang graduate."),
        ("I can send and receive messages in direct and group conversations.",
         "Maaari akong magpadala at makatanggap ng messages sa direct at group conversations."),
        ("I can attach a supported image or file to a message.",
         "Maaari akong mag-attach ng supported image o file sa isang message."),
        ("I can see unread counts, online or last-active information, and delivery or read status where available.",
         "Maaari kong makita ang unread counts, online o last-active information, at delivery o read status kapag available."),
        ("I can delete a message that I sent.",
         "Maaari kong i-delete ang isang message na ako ang nagpadala."),
        ("In a group conversation, I can view group information and use add-member or leave-group actions when permitted.",
         "Sa group conversation, maaari kong makita ang group information at gamitin ang add-member o leave-group actions kapag permitted."),
    ],
    [
        ("I can view another graduate's community profile and published posts when the community feature is enabled.",
         "Maaari kong makita ang community profile at published posts ng ibang graduate kapag enabled ang community feature."),
        ("I can update my personal and contact details.",
         "Maaari kong i-update ang aking personal at contact details."),
        ("I can update my employment and career details.",
         "Maaari kong i-update ang aking employment at career details."),
        ("I can update my education and graduation details.",
         "Maaari kong i-update ang aking education at graduation details."),
        ("I can upload, replace, or remove my profile photo.",
         "Maaari akong mag-upload, magpalit, o mag-remove ng aking profile photo."),
        ("I can upload, replace, or remove my cover photo.",
         "Maaari akong mag-upload, magpalit, o mag-remove ng aking cover photo."),
        ("I can change my password from the profile settings.",
         "Maaari kong palitan ang aking password mula sa profile settings."),
        ("I can view Graduate Portal notifications when the notification feature is enabled.",
         "Maaari kong makita ang Graduate Portal notifications kapag enabled ang notification feature."),
    ],
]

# The source template intentionally allocates 8, 6, 10, 6, and 10 statements
# across its five role-function tables. Repartition the validated 40 statements
# without changing the document's table structure.
ALUMNI_GRADUATE_GROUPS = [
    ALUMNI_GRADUATE_GROUPS[0],
    ALUMNI_GRADUATE_GROUPS[1][:6],
    ALUMNI_GRADUATE_GROUPS[1][6:] + ALUMNI_GRADUATE_GROUPS[2],
    ALUMNI_GRADUATE_GROUPS[3][:6],
    ALUMNI_GRADUATE_GROUPS[3][6:] + ALUMNI_GRADUATE_GROUPS[4],
]


ALUMNI_PRESIDENT_GROUPS = [
    [
        ("I can view account requests that are pending Alumni Admin verification.",
         "Maaari kong makita ang account requests na pending para sa Alumni Admin verification."),
        ("I can search account requests and filter them as Pending, Approved, or Rejected.",
         "Maaari akong maghanap ng account requests at i-filter ang mga ito bilang Pending, Approved, o Rejected."),
        ("I can open an account request to review its identity, program, graduation year, contact, and survey-submission information.",
         "Maaari kong buksan ang account request upang suriin ang identity, program, graduation year, contact, at survey-submission information nito."),
        ("I can review whether an account request matches a linked alumni-registry record.",
         "Maaari kong suriin kung tumutugma ang account request sa linked alumni-registry record."),
        ("I can approve a pending Graduate Portal account request.",
         "Maaari kong aprubahan ang pending Graduate Portal account request."),
        ("I can reject a pending Graduate Portal account request and provide a rejection reason.",
         "Maaari kong tanggihan ang pending Graduate Portal account request at magbigay ng rejection reason."),
        ("I can view who reviewed an account request and when it was reviewed.",
         "Maaari kong makita kung sino ang nag-review ng account request at kung kailan ito na-review."),
        ("The approved or rejected status is shown after an account request is reviewed.",
         "Ipinapakita ang approved o rejected status pagkatapos ma-review ang account request."),
    ],
    [
        ("I can import an alumni list and preview valid, duplicate, invalid, and ignored rows before confirming.",
         "Maaari akong mag-import ng alumni list at i-preview ang valid, duplicate, invalid, at ignored rows bago magkumpirma."),
        ("I can choose how duplicate alumni rows are handled during import.",
         "Maaari kong piliin kung paano hahawakan ang duplicate alumni rows sa import."),
        ("I can download a list of import errors when errors are present.",
         "Maaari akong mag-download ng listahan ng import errors kapag may errors."),
        ("I can search registry records and filter them by course, batch, registration status, and survey-answer status.",
         "Maaari akong maghanap ng registry records at i-filter ang mga ito ayon sa course, batch, registration status, at survey-answer status."),
        ("I can open an alumni registry record to view its registry and linked-account details.",
         "Maaari kong buksan ang alumni registry record upang makita ang registry at linked-account details nito."),
        ("I can edit registry information and linked portal-account details when an account is linked.",
         "Maaari kong i-edit ang registry information at linked portal-account details kapag may linked account."),
        ("I can mark an alumni registry record as Verified or Inactive.",
         "Maaari kong markahan ang alumni registry record bilang Verified o Inactive."),
        ("I can export selected registry data in an available file format and scope.",
         "Maaari kong i-export ang napiling registry data gamit ang available na file format at scope."),
    ],
    [
        ("I can switch between the Registered Alumni list and Alumni Archive.",
         "Maaari akong lumipat sa pagitan ng Registered Alumni list at Alumni Archive."),
        ("I can archive an active alumni registry record after confirmation.",
         "Maaari kong i-archive ang active alumni registry record pagkatapos ng confirmation."),
        ("Archiving moves the record out of the active alumni list without permanently deleting related data.",
         "Inililipat ng archiving ang record mula sa active alumni list nang hindi permanenteng dini-delete ang related data."),
        ("I can view when an archived alumni record was archived and who archived it.",
         "Maaari kong makita kung kailan na-archive ang alumni record at kung sino ang nag-archive nito."),
        ("I can search and filter records while viewing the Alumni Archive.",
         "Maaari akong mag-search at mag-filter ng records habang nasa Alumni Archive."),
        ("I can restore an archived alumni registry record.",
         "Maaari kong i-restore ang archived alumni registry record."),
        ("I can permanently delete an alumni registry record only from the Alumni Archive.",
         "Maaari kong permanenteng i-delete ang alumni registry record mula lamang sa Alumni Archive."),
        ("The permanent-delete warning explains that the linked graduate account is not deleted.",
         "Ipinapaliwanag ng permanent-delete warning na hindi dini-delete ang linked graduate account."),
    ],
    [
        ("I can create an announcement with a title, content, publication status, and optional images.",
         "Maaari akong gumawa ng announcement na may title, content, publication status, at optional images."),
        ("I can filter announcements as Published, Draft, or Archived and edit or delete an announcement.",
         "Maaari kong i-filter ang announcements bilang Published, Draft, o Archived at i-edit o i-delete ang announcement."),
        ("I can view counts for pending, resolved, and dismissed forum reports.",
         "Maaari kong makita ang counts para sa pending, resolved, at dismissed forum reports."),
        ("I can search forum reports and filter them by report type, reason, and status.",
         "Maaari akong maghanap ng forum reports at i-filter ang mga ito ayon sa report type, reason, at status."),
        ("I can open a report to review the reported post or comment and its report details.",
         "Maaari kong buksan ang report upang suriin ang reported post o comment at report details nito."),
        ("I can hide reported content or restore content that was hidden.",
         "Maaari kong itago ang reported content o i-restore ang content na itinago."),
        ("I can resolve a pending forum report.",
         "Maaari kong i-resolve ang pending forum report."),
        ("I can dismiss a pending forum report.",
         "Maaari kong i-dismiss ang pending forum report."),
    ],
    [
        ("I can view the Pending, Approved, and Declined counts for submitted job posts.",
         "Maaari kong makita ang Pending, Approved, at Declined counts para sa submitted job posts."),
        ("I can search job posts by graduate, program, company, skills, or job title.",
         "Maaari akong maghanap ng job posts ayon sa graduate, program, company, skills, o job title."),
        ("I can filter job posts by Pending, Approved, or Declined status.",
         "Maaari kong i-filter ang job posts ayon sa Pending, Approved, o Declined status."),
        ("I can open a job post to review its poster, job details, program fit, deadline, and application information.",
         "Maaari kong buksan ang job post upang suriin ang poster, job details, program fit, deadline, at application information nito."),
        ("I can enter optional review notes before deciding on a pending job post.",
         "Maaari akong maglagay ng optional review notes bago magdesisyon sa pending job post."),
        ("I can approve a pending job post so that it becomes visible in the Graduate Portal.",
         "Maaari kong aprubahan ang pending job post upang maging visible ito sa Graduate Portal."),
        ("I can decline a pending job post so that it remains hidden from graduates.",
         "Maaari kong tanggihan ang pending job post upang manatili itong hidden sa mga graduate."),
        ("I can view the review status, notes, reviewer, and review time for a processed job post.",
         "Maaari kong makita ang review status, notes, reviewer, at review time para sa processed job post."),
    ],
]


SPECS = [
    {
        "source": "GradTrack_Questionnaire_Research_Coordinator_Revised.docx",
        "output": "GradTrack_Questionnaire_Research_Coordinator_System_Aligned.docx",
        "paragraphs": {
            3: "Main responsibility / Pangunahing responsibilidad: Read-only graduate tracer analytics / Read-only na graduate tracer analytics",
            11: "PART II. ASSESSMENT OF RESEARCH DASHBOARD INFORMATION / BAHAGI II. PAGSUSURI SA IMPORMASYON NG RESEARCH DASHBOARD",
            12: "This section covers the implemented dashboard information available to the Research Coordinator. / Saklaw ng bahaging ito ang implemented dashboard information na available sa Research Coordinator.",
            13: "1.1 Dashboard Access and Key Metrics / Dashboard Access at Mahahalagang Metrics",
            14: "1.2 Survey Participation Summary / Buod ng Survey Participation",
            15: "1.3 Employability by Program / Employability Ayon sa Programa",
            16: "1.4 Employment Trends / Employment Trends",
            17: "1.5 Job Alignment Distribution / Job Alignment Distribution",
            18: "PART III. ASSESSMENT OF RESEARCH COORDINATOR FUNCTIONS / BAHAGI III. PAGSUSURI SA RESEARCH COORDINATOR FUNCTIONS",
            19: "This section covers the implemented read-only analytics and account controls available to the Research Coordinator. / Saklaw ng bahaging ito ang implemented read-only analytics at account controls na available sa Research Coordinator.",
            20: "2.1 Selected Survey Snapshot / Selected Survey Snapshot",
            21: "2.2 Employment and Alignment Bases / Batayan ng Employment at Alignment",
            22: "2.3 Interactive Chart Information / Interactive na Impormasyon ng Charts",
            23: "2.4 Visual Analytics and Empty-Data Messages / Visual Analytics at Empty-Data Messages",
            24: "2.5 Read-Only Access and Account Controls / Read-Only Access at Account Controls",
        },
        "groups": RESEARCH_GROUPS,
        "start_table": 2,
    },
    {
        "source": "GradTrack_Questionnaire_Dean_Revised.docx",
        "output": "GradTrack_Questionnaire_Dean_System_Aligned.docx",
        "paragraphs": {
            3: "Main responsibility / Pangunahing responsibilidad: Program-scoped survey participation monitoring and reminder emails / Program-scoped na survey participation monitoring at reminder emails",
            12: "2.1 Program Scope and Participation Summary / Program Scope at Participation Summary",
            13: "2.2 Search and Participation Filters / Search at Participation Filters",
            14: "2.3 Participation Records and Selection / Participation Records at Pagpili",
            15: "2.4 Read-Only Submitted Answers / Read-Only na Submitted Answers",
            16: "2.5 Reminder Email Workflow / Proseso ng Reminder Email",
        },
        "groups": DEAN_GROUPS,
        "start_table": 2,
    },
    {
        "source": "GradTrack_Questionnaire_Registrar_Revised.docx",
        "output": "GradTrack_Questionnaire_Registrar_System_Aligned.docx",
        "paragraphs": {
            3: "Main responsibility / Pangunahing responsibilidad: Official graduate list entry, import, filtering, and record archiving / Pag-entry, pag-import, pag-filter, at pag-archive ng official graduate list",
            12: "2.1 Manual Graduate Record Addition / Manwal na Pagdagdag ng Graduate Record",
            13: "2.2 XLSX and CSV Graduate Import / Pag-import ng Graduate XLSX at CSV",
            14: "2.3 Search, Department, and Year Filters / Search, Department, at Year Filters",
            15: "2.4 Selection and Archiving / Pagpili at Pag-archive",
            16: "2.5 Registrar Archive, Restore, and Permanent Delete / Registrar Archive, Restore, at Permanent Delete",
        },
        "groups": REGISTRAR_GROUPS,
        "start_table": 2,
    },
    {
        "source": "GradTrack_Questionnaire_Super_Admin_Revised.docx",
        "output": "GradTrack_Questionnaire_Super_Admin_System_Aligned.docx",
        "paragraphs": {
            3: "Main responsibility / Pangunahing responsibilidad: Staff accounts, reminders, audit, backup, and system configuration / Staff accounts, reminders, audit, backup, at system configuration",
            12: "2.1 Staff Accounts and Roles / Staff Accounts at Roles",
            13: "2.2 Survey Reminder Monitoring and Configuration / Survey Reminder Monitoring at Configuration",
            14: "2.3 Audit Trail and Database Backup / Audit Trail at Database Backup",
            15: "2.4 General, Branding, and Login Settings / General, Branding, at Login Settings",
            16: "2.5 Feature, Survey, and Community Settings / Feature, Survey, at Community Settings",
            17: "2.6 Maintenance Mode / Maintenance Mode",
            18: "2.7 Public Website Content / Public Website Content",
            19: "2.8 Settings Workflow and Profile / Settings Workflow at Profile",
        },
        "groups": SUPER_ADMIN_GROUPS,
        "start_table": 2,
    },
    {
        "source": "GradTrack_Questionnaire_Alumni_Graduate_Revised.docx",
        "output": "GradTrack_Questionnaire_Alumni_Graduate_System_Aligned.docx",
        "paragraphs": {
            3: "Main responsibility / Pangunahing responsibilidad: Tracer survey participation and alumni services through the Graduate Portal / Pagsagot sa tracer survey at paggamit ng alumni services sa Graduate Portal",
            12: "2.1 Survey Verification, Completion, and Account Request / Survey Verification, Completion, at Account Request",
            13: "2.2 Announcements, Dashboard, and Job Support / Announcements, Dashboard, at Job Support",
            14: "2.3 Job Posting and Community Forum / Job Posting at Community Forum",
            15: "2.4 Direct and Group Messaging / Direct at Group Messaging",
            16: "2.5 Conversation Controls, Alumni Profiles, and Settings / Conversation Controls, Alumni Profiles, at Settings",
        },
        "groups": ALUMNI_GRADUATE_GROUPS,
        "start_table": 2,
    },
    {
        "source": "GradTrack_Questionnaire_Alumni_President_Revised.docx",
        "output": "GradTrack_Questionnaire_Alumni_President_System_Aligned.docx",
        "paragraphs": {
            2: "Respondent: Alumni President (Alumni Admin)",
            3: "Main responsibility / Pangunahing responsibilidad: Alumni account verification, registry, announcements, forum reports, and job approvals / Alumni account verification, registry, announcements, forum reports, at job approvals",
            12: "2.1 Graduate Portal Account Verification / Graduate Portal Account Verification",
            13: "2.2 Alumni Registry Import, Review, and Export / Alumni Registry Import, Review, at Export",
            14: "2.3 Registry Archive, Restore, and Permanent Delete / Registry Archive, Restore, at Permanent Delete",
            15: "2.4 Announcements and Forum Report Moderation / Announcements at Forum Report Moderation",
            16: "2.5 Job Post Review and Approval / Job Post Review at Approval",
        },
        "table_role": "Alumni President (Alumni Admin)",
        "groups": ALUMNI_PRESIDENT_GROUPS,
        "start_table": 2,
    },
]


def set_question_groups(document: Document, start_table: int, groups: list[list[tuple[str, str]]]) -> None:
    number = 1
    for offset, questions in enumerate(groups):
        table = document.tables[start_table + offset]
        expected_rows = len(questions) + 1
        if len(table.rows) != expected_rows:
            raise ValueError(
                f"Table {start_table + offset} has {len(table.rows)} rows; expected {expected_rows}."
            )

        for row_index, (english, tagalog) in enumerate(questions, start=1):
            table.cell(row_index, 0).text = str(number)
            table.cell(row_index, 1).text = f"English: {english}\nTagalog: {tagalog}"
            number += 1

    if number != 41:
        raise ValueError(f"Expected 40 role-function questions, found {number - 1}.")


def add_row_property(row, property_name: str) -> None:
    row_properties = row._tr.get_or_add_trPr()
    qualified_name = qn(f"w:{property_name}")
    if row_properties.find(qualified_name) is None:
        row_properties.append(OxmlElement(f"w:{property_name}"))


def improve_print_layout(document: Document) -> None:
    for paragraph in document.paragraphs:
        if paragraph.style.name in {"Heading 1", "Heading 2"}:
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True

    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            add_row_property(row, "cantSplit")
            if row_index == 0:
                add_row_property(row, "tblHeader")


def validate_saved_document(path: Path, start_table: int, groups: list[list[tuple[str, str]]]) -> None:
    document = Document(path)
    expected = list(range(1, 41))
    actual: list[int] = []
    for offset, questions in enumerate(groups):
        table = document.tables[start_table + offset]
        if len(table.rows) != len(questions) + 1:
            raise ValueError(f"Saved table {start_table + offset} has an unexpected row count.")
        actual.extend(int(table.cell(row_index, 0).text) for row_index in range(1, len(table.rows)))
    if actual != expected:
        raise ValueError(f"Saved question numbering is invalid: {actual}")


def revise(spec: dict) -> Path:
    source = SOURCE_DIR / spec["source"]
    if not source.exists():
        raise FileNotFoundError(source)

    document = Document(source)
    document.paragraphs[6].text = INTRO_EN
    document.paragraphs[7].text = INTRO_TL

    for paragraph_index, value in spec.get("paragraphs", {}).items():
        document.paragraphs[paragraph_index].text = value

    if "table_role" in spec:
        document.tables[0].cell(0, 1).text = spec["table_role"]

    set_question_groups(document, spec["start_table"], spec["groups"])
    improve_print_layout(document)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output = OUTPUT_DIR / spec["output"]
    document.save(output)
    validate_saved_document(output, spec["start_table"], spec["groups"])
    return output


if __name__ == "__main__":
    for questionnaire_spec in SPECS:
        print(revise(questionnaire_spec))
