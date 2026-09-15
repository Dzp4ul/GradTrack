from pathlib import Path
import sys
from zipfile import ZipFile

from docx import Document


DOCS = [
    Path(r"C:\Users\celvn\Downloads\GradTrack_Questionnaire_Research_Coordinator_Revised.docx"),
    Path(r"C:\Users\celvn\Downloads\GradTrack_Questionnaire_Dean_Revised.docx"),
    Path(r"C:\Users\celvn\Downloads\GradTrack_Questionnaire_Registrar_Revised.docx"),
    Path(r"C:\Users\celvn\Downloads\GradTrack_Questionnaire_Super_Admin_Revised.docx"),
    Path(r"C:\Users\celvn\Downloads\GradTrack_Questionnaire_Alumni_Graduate_Revised.docx"),
    Path(r"C:\Users\celvn\Downloads\GradTrack_Questionnaire_Alumni_President_Revised.docx"),
]

if len(sys.argv) > 1:
    selected = sys.argv[1].lower()
    DOCS = [path for path in DOCS if selected in path.name.lower()]


def text(value: str) -> str:
    return value.replace("\n", "\\n")


for path in DOCS:
    document = Document(path)
    print(f"\n{'=' * 100}\n{path.name}")
    print(f"sections={len(document.sections)} paragraphs={len(document.paragraphs)} tables={len(document.tables)}")
    print("PARAGRAPHS")
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.text.strip():
            print(f"P{index:03d} [{paragraph.style.name}] {text(paragraph.text)}")
    print("TABLES")
    for table_index, table in enumerate(document.tables):
        print(f"T{table_index:02d} rows={len(table.rows)} cols={len(table.columns)} style={table.style.name if table.style else ''}")
        for row_index, row in enumerate(table.rows):
            cells = [text(cell.text) for cell in row.cells]
            print(f"  R{row_index:02d}: " + " | ".join(cells))
    for section_index, section in enumerate(document.sections):
        print(f"SECTION {section_index}: {section.page_width}x{section.page_height}; margins={section.top_margin},{section.right_margin},{section.bottom_margin},{section.left_margin}")
        print("  HEADER:", " | ".join(text(p.text) for p in section.header.paragraphs if p.text.strip()))
        print("  FOOTER:", " | ".join(text(p.text) for p in section.footer.paragraphs if p.text.strip()))
    with ZipFile(path) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        print("MEDIA:", media)
