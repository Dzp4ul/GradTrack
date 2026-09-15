from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn


OUTPUT_DIR = Path(__file__).resolve().parents[1] / "docs" / "revised_questionnaires"
BANNED_CLAIMS = [
    "filter participation records by survey",
    "select the survey whose participation",
    "update graduate information when changes are reported",
    "distribute the tracer survey to the intended graduates",
    "prepare graduate tracer study reports from collected responses",
    "restore a database backup from the system",
    "update my account name, email, profile image",
    "super admin profile name, email, profile image",
]


def validate(path: Path) -> None:
    if ZipFile(path).testzip() is not None:
        raise ValueError(f"Corrupt DOCX package: {path.name}")

    document = Document(path)
    all_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    ).lower()
    found_banned = [claim for claim in BANNED_CLAIMS if claim in all_text]
    if found_banned:
        raise ValueError(f"Unsupported claim in {path.name}: {found_banned}")

    item_count = sum(max(0, len(table.rows) - 1) for table in document.tables[2:])
    if item_count != 60:
        raise ValueError(f"Expected 60 rated items in {path.name}; found {item_count}")

    for table in document.tables:
        if table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is None:
            raise ValueError(f"Missing repeating header row in {path.name}")
        for row in table.rows:
            if row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is None:
                raise ValueError(f"Splittable table row in {path.name}")


if __name__ == "__main__":
    paths = sorted(OUTPUT_DIR.glob("*_System_Aligned.docx"))
    if len(paths) != 6:
        raise ValueError(f"Expected 6 revised questionnaires; found {len(paths)}")
    for questionnaire in paths:
        validate(questionnaire)
        print(f"{questionnaire.name}: OK")
